"""
test_system.py — Automated System Testing Script
-------------------------------------------------
Reads dataset, runs detector on every sample, saves results to Excel.

DATASET FORMAT EXPECTED:
  Column A (#)       — row number (ignored)
  Column B (Type)    — true label: "Safe" or "Not Safe"
  Column C (Source)  — source info (ignored)
  Column D (Amount)  — amount info (ignored)
  Column E (Content) — the actual URL or message to test ← KEY FIX

SETTINGS
  DATASET_FILE : path to your Excel dataset
  USE_LLM      : True  = include Ollama (accurate but slow ~30-60 min)
                 False = rule+ML only (fast, no LLM penalty applied)
"""

import openpyxl
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from datetime import datetime

from detector import analyze_for_ui

# ── SETTINGS ──────────────────────────────────────────────────────────────────
DATASET_FILE = "dataset (2).xlsx"
USE_LLM      = True   # set True to include Ollama (slower but more accurate)


# ── STYLES ────────────────────────────────────────────────────────────────────
HDR_FONT  = Font(bold=True, color="FFFFFF", size=11, name="Arial")
HDR_FILL  = PatternFill("solid", fgColor="1F4E79")
ALT_FILL  = PatternFill("solid", fgColor="F2F2F2")
OK_FILL   = PatternFill("solid", fgColor="C6EFCE")
ERR_FILL  = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFE699")

def thin_border():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)

def hdr_cell(cell, value):
    cell.value     = value
    cell.font      = HDR_FONT
    cell.fill      = HDR_FILL
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
    cell.border    = thin_border()

def data_cell(cell, value, alt_row=False, fill_override=None):
    cell.value     = value
    cell.font      = Font(name="Arial", size=10)
    cell.fill      = fill_override if fill_override else (ALT_FILL if alt_row else PatternFill("solid", fgColor="FFFFFF"))
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.border    = thin_border()


# ============================================================
# STEP 1 — LOAD DATASET
# ============================================================
def load_dataset(filepath):
    """
    Reads dataset using pandas for reliable column mapping.
    Expected columns: #, Type, Source, Amount, Content
    """
    print(f"📂 Loading dataset: {filepath}")
    df = pd.read_excel(filepath)

    # Validate required columns exist
    required = ["Type", "Content"]
    missing  = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"Dataset is missing columns: {missing}. Found: {df.columns.tolist()}")

    label_map = {
        "safe":      "Safe",
        "not safe":  "Suspicious",
        "notsafe":   "Suspicious",
        "suspicious":"Suspicious",
        "phishing":  "Suspicious",
        "malicious": "Suspicious",
        "unsafe":    "Suspicious",
    }

    data    = []
    skipped = 0

    for _, row in df.iterrows():
        content   = str(row["Content"]).strip() if pd.notna(row["Content"]) else ""
        raw_label = str(row["Type"]).strip()    if pd.notna(row["Type"])    else ""

        if not content or not raw_label:
            skipped += 1
            continue

        true_label = label_map.get(raw_label.lower(), None)
        if true_label is None:
            print(f"  ⚠️  Skipped unrecognised label: '{raw_label}'")
            skipped += 1
            continue

        data.append({"input": content, "true_label": true_label})

    print(f"  ✅ Loaded {len(data)} samples ({skipped} skipped)")

    # Show label breakdown
    safe_count = sum(1 for d in data if d["true_label"] == "Safe")
    susp_count = sum(1 for d in data if d["true_label"] == "Suspicious")
    print(f"  📊 Safe: {safe_count} | Suspicious: {susp_count}")
    return data


# ============================================================
# STEP 2 — RUN DETECTION
# ============================================================
def run_tests(data):
    results = []
    total   = len(data)

    for i, sample in enumerate(data, 1):
        print(f"  🔍 [{i}/{total}] {sample['input'][:70]}...")

        try:
            result = analyze_for_ui(sample["input"], use_llm=USE_LLM)

            # Normalize predicted label → Safe or Suspicious
            raw_predicted = result["final"]["label"].strip().lower()
            if any(w in raw_predicted for w in ["suspicious", "phishing", "malicious", "unsafe", "not safe"]):
                predicted_binary = "Suspicious"
            elif "safe" in raw_predicted:
                predicted_binary = "Safe"
            else:
                predicted_binary = result["final"]["label"].strip().capitalize()

            true_label = sample["true_label"]
            correct    = predicted_binary == true_label

            results.append({
                "Input":            sample["input"],
                "Type":             result["input"]["type"],
                "True Label":       true_label,
                "Predicted":        result["final"]["label"],
                "Predicted Binary": predicted_binary,
                "Correct":          "✅" if correct else "❌",
                "Final Score":      result["final"]["score"],
                "Rule Risk":        result["breakdown"]["rule_risk"],
                "ML Risk":          result["breakdown"]["ml_risk"],
                "LLM Score":        result["breakdown"]["llm_score"],
                "Red Flags":        ", ".join(result["final"]["red_flags"]),
                "Top Reason":       result["final"]["reasons"][0] if result["final"]["reasons"] else "",
            })

        except Exception as e:
            results.append({
                "Input":            sample["input"],
                "Type":             "ERROR",
                "True Label":       sample["true_label"],
                "Predicted":        "ERROR",
                "Predicted Binary": "ERROR",
                "Correct":          "❌",
                "Final Score":      0,
                "Rule Risk":        0,
                "ML Risk":          0,
                "LLM Score":        0,
                "Red Flags":        "",
                "Top Reason":       f"Error: {e}",
            })

    return results


# ============================================================
# STEP 3 — CALCULATE METRICS
# ============================================================
def calculate_metrics(results):
    tp = sum(1 for r in results if r["True Label"] == "Suspicious" and r["Predicted Binary"] == "Suspicious")
    tn = sum(1 for r in results if r["True Label"] == "Safe"        and r["Predicted Binary"] == "Safe")
    fp = sum(1 for r in results if r["True Label"] == "Safe"        and r["Predicted Binary"] == "Suspicious")
    fn = sum(1 for r in results if r["True Label"] == "Suspicious"  and r["Predicted Binary"] == "Safe")

    total   = len(results)
    correct = tp + tn

    accuracy  = round((tp + tn) / total * 100, 2) if total > 0       else 0
    precision = round(tp / (tp + fp) * 100, 2)    if (tp + fp) > 0   else 0
    recall    = round(tp / (tp + fn) * 100, 2)    if (tp + fn) > 0   else 0
    f1        = round(2 * precision * recall / (precision + recall), 2) if (precision + recall) > 0 else 0

    return {
        "Total Samples":          total,
        "Correct Predictions":    correct,
        "Wrong Predictions":      total - correct,
        "True Positive (TP)":     tp,
        "True Negative (TN)":     tn,
        "False Positive (FP)":    fp,
        "False Negative (FN)":    fn,
        "Accuracy (%)":           accuracy,
        "Precision (%)":          precision,
        "Recall (%)":             recall,
        "F1 Score (%)":           f1,
    }


# ============================================================
# STEP 4 — GENERATE CHARTS
# ============================================================
def generate_metrics_chart(metrics, path="metrics_chart.png"):
    labels = ["Accuracy", "Precision", "Recall", "F1 Score"]
    values = [metrics["Accuracy (%)"], metrics["Precision (%)"], metrics["Recall (%)"], metrics["F1 Score (%)"]]
    colors = ["#1F4E79", "#2E75B6", "#5BA3D9", "#9DC3E6"]

    plt.figure(figsize=(8, 5))
    bars = plt.bar(labels, values, color=colors)
    plt.ylim(0, 110)
    plt.ylabel("Score (%)")
    plt.title("System Performance Metrics")
    for bar, val in zip(bars, values):
        plt.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 1.5,
                 f"{val}%", ha="center", fontsize=11, fontweight="bold")
    plt.tight_layout()
    plt.savefig(path, dpi=150)
    plt.close()
    print(f"  📊 Saved: {path}")


def generate_confusion_matrix(metrics, path="confusion_matrix.png"):
    tp = metrics["True Positive (TP)"]
    tn = metrics["True Negative (TN)"]
    fp = metrics["False Positive (FP)"]
    fn = metrics["False Negative (FN)"]

    matrix = [[tn, fp], [fn, tp]]
    labels = ["Safe", "Suspicious"]
    max_val = max(tp, tn, fp, fn) if max(tp, tn, fp, fn) > 0 else 1

    fig, ax = plt.subplots(figsize=(6, 5))
    im = ax.imshow(matrix, interpolation="nearest", cmap="Blues")
    plt.colorbar(im)
    ax.set_xticks([0, 1]); ax.set_yticks([0, 1])
    ax.set_xticklabels(labels); ax.set_yticklabels(labels)
    ax.set_xlabel("Predicted Label"); ax.set_ylabel("True Label")
    ax.set_title("Confusion Matrix")
    for i in range(2):
        for j in range(2):
            ax.text(j, i, str(matrix[i][j]), ha="center", va="center",
                    fontsize=14, fontweight="bold",
                    color="white" if matrix[i][j] > max_val / 2 else "black")
    plt.tight_layout()
    plt.savefig(path, dpi=150)
    plt.close()
    print(f"  📊 Saved: {path}")


# ============================================================
# STEP 5 — SAVE TO EXCEL
# ============================================================
def save_to_excel(results, metrics):
    wb  = openpyxl.Workbook()

    # ── Sheet 1: Test Results ──────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Test Results"

    headers = ["No.", "Input", "Type", "True Label", "Predicted",
               "Correct", "Final Score", "Rule Risk", "ML Risk",
               "LLM Score", "Red Flags", "Top Reason"]

    for col, h in enumerate(headers, 1):
        hdr_cell(ws1.cell(row=1, column=col), h)

    for row_num, r in enumerate(results, 2):
        alt = row_num % 2 == 0
        ws1.cell(row=row_num, column=1).value  = row_num - 1
        ws1.cell(row=row_num, column=1).border = thin_border()
        ws1.cell(row=row_num, column=1).alignment = Alignment(horizontal="center")

        values = [r["Input"], r["Type"], r["True Label"], r["Predicted"],
                  r["Correct"], r["Final Score"], r["Rule Risk"], r["ML Risk"],
                  r["LLM Score"], r["Red Flags"], r["Top Reason"]]

        for col, val in enumerate(values, 2):
            cell = ws1.cell(row=row_num, column=col)
            if col == 6:  # Correct column
                fill = OK_FILL if val == "✅" else ERR_FILL
            else:
                fill = None
            data_cell(cell, val, alt_row=alt, fill_override=fill)

    col_widths = {"A":5,"B":55,"C":10,"D":14,"E":14,
                  "F":9,"G":12,"H":12,"I":10,"J":10,"K":35,"L":50}
    for col, w in col_widths.items():
        ws1.column_dimensions[col].width = w
    ws1.freeze_panes = "A2"
    ws1.row_dimensions[1].height = 20

    # ── Sheet 2: Performance Metrics ──────────────────────────────────────
    ws2 = wb.create_sheet("Performance Metrics")
    hdr_cell(ws2.cell(row=1, column=1), "Metric")
    hdr_cell(ws2.cell(row=1, column=2), "Value")

    for row_num, (metric, value) in enumerate(metrics.items(), 2):
        alt = row_num % 2 == 0
        data_cell(ws2.cell(row=row_num, column=1), metric, alt_row=alt)
        cell = ws2.cell(row=row_num, column=2)

        if "%" in str(metric) and isinstance(value, float):
            fill = OK_FILL if value >= 80 else (WARN_FILL if value >= 60 else ERR_FILL)
        else:
            fill = None
        data_cell(cell, value, alt_row=alt, fill_override=fill)
        cell.alignment = Alignment(horizontal="center")

    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 15

    # ── Sheet 3: Metrics Chart ─────────────────────────────────────────────
    generate_metrics_chart(metrics)
    ws3 = wb.create_sheet("Metrics Chart")
    img1 = XLImage("metrics_chart.png")
    img1.anchor = "A1"
    ws3.add_image(img1)

    # ── Sheet 4: Confusion Matrix ──────────────────────────────────────────
    generate_confusion_matrix(metrics)
    ws4 = wb.create_sheet("Confusion Matrix")
    img2 = XLImage("confusion_matrix.png")
    img2.anchor = "A1"
    ws4.add_image(img2)

    # ── Save ──────────────────────────────────────────────────────────────
    filename = f"System_Test_Results_{datetime.today().strftime('%Y-%m-%d')}.xlsx"
    wb.save(filename)
    print(f"\n  ✅ Saved: {filename}")
    return filename


# ============================================================
# RUN
# ============================================================
if __name__ == "__main__":
    print("=" * 60)
    print("  🚀 AUTOMATED SYSTEM TESTING")
    print(f"  Dataset : {DATASET_FILE}")
    print(f"  LLM     : {'ENABLED (Ollama)' if USE_LLM else 'DISABLED (Rule+ML only)'}")
    print("=" * 60 + "\n")

    print("STEP 1 — Loading dataset...")
    data = load_dataset(DATASET_FILE)

    print("\nSTEP 2 — Running detection on all samples...")
    results = run_tests(data)

    print("\nSTEP 3 — Calculating metrics...")
    metrics = calculate_metrics(results)

    print("\n📋 Performance Summary:")
    print("─" * 40)
    for k, v in metrics.items():
        print(f"  {k:<28}: {v}")
    print("─" * 40)

    print("\nSTEP 4 — Saving to Excel...")
    save_to_excel(results, metrics)

    print("\n🎉 Testing complete!")